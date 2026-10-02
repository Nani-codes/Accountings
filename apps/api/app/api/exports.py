from __future__ import annotations

import io
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import AiFinding, Client, Invoice, ReconciliationResult, ReconciliationRun, User
from app.services.auth import get_current_user

router = APIRouter(tags=["exports"])


def _get_run(db: Session, run_id: UUID, org_id: UUID) -> ReconciliationRun:
    run = (
        db.query(ReconciliationRun)
        .filter(
            ReconciliationRun.id == run_id,
            ReconciliationRun.organization_id == org_id,
        )
        .first()
    )
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
    return run


def _accepted_findings(db: Session, run_id: UUID) -> list[AiFinding]:
    return (
        db.query(AiFinding)
        .filter(
            AiFinding.run_id == run_id,
            AiFinding.status.in_(("accepted", "edited")),
        )
        .all()
    )


@router.get("/runs/{run_id}/export.xlsx")
def export_excel(
    run_id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    run = _get_run(db, run_id, user.organization_id)
    client = db.query(Client).filter(Client.id == run.client_id).one()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Client", client.name])
    summary.append(["Period", run.period])
    summary.append(["Status", run.status])
    for k, v in (run.summary or {}).items():
        summary.append([k, str(v)])

    matches = wb.create_sheet("Matches")
    matches.append(["match_status", "difference_type", "difference_amount", "score"])
    for row in db.query(ReconciliationResult).filter(ReconciliationResult.run_id == run.id):
        matches.append(
            [
                row.match_status,
                row.difference_type,
                float(row.difference_amount) if row.difference_amount is not None else None,
                float(row.match_score) if row.match_score is not None else None,
            ]
        )

    exceptions = wb.create_sheet("Exceptions")
    exceptions.append(["match_status", "difference_type", "difference_amount"])
    for row in db.query(ReconciliationResult).filter(
        ReconciliationResult.run_id == run.id,
        ReconciliationResult.match_status != "exact_match",
    ):
        exceptions.append(
            [
                row.match_status,
                row.difference_type,
                float(row.difference_amount) if row.difference_amount is not None else None,
            ]
        )

    findings_sheet = wb.create_sheet("Findings")
    findings_sheet.append(["title", "category", "severity", "status", "recommended_action"])
    for f in _accepted_findings(db, run.id):
        findings_sheet.append([f.title, f.category, f.severity, f.status, f.recommended_action])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="working-paper-{run.period}.xlsx"'},
    )


@router.get("/runs/{run_id}/export.pdf")
def export_pdf(
    run_id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    run = _get_run(db, run_id, user.organization_id)
    client = db.query(Client).filter(Client.id == run.client_id).one()
    findings = _accepted_findings(db, run.id)

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "GST RECONCILIATION WORKING PAPER")
    y -= 24
    c.setFont("Helvetica", 11)
    c.drawString(50, y, f"Client: {client.name}")
    y -= 16
    c.drawString(50, y, f"Period: {run.period}")
    y -= 16
    summary = run.summary or {}
    c.drawString(50, y, f"Matched: {summary.get('matched', '-')}")
    y -= 16
    c.drawString(50, y, f"Exceptions: {summary.get('exceptions', '-')}")
    y -= 28
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "Findings (accepted/edited)")
    y -= 18
    c.setFont("Helvetica", 10)
    for f in findings:
        if y < 80:
            c.showPage()
            y = height - 50
            c.setFont("Helvetica", 10)
        c.drawString(50, y, f"- {f.title} [{f.category}]")
        y -= 14
    y -= 20
    c.drawString(50, y, "Prepared by: AI Assistant")
    y -= 16
    c.drawString(50, y, "Reviewed by: ________________")
    c.showPage()
    c.save()
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="working-paper-{run.period}.pdf"'},
    )


@router.post("/runs/{run_id}/working-paper")
def generate_working_paper(
    run_id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Return structured working-paper payload (AI narrative wired in Task 10)."""
    run = _get_run(db, run_id, user.organization_id)
    client = db.query(Client).filter(Client.id == run.client_id).one()
    findings = _accepted_findings(db, run.id)
    return {
        "client": client.name,
        "period": run.period,
        "summary": run.summary or {},
        "findings": [
            {
                "title": f.title,
                "category": f.category,
                "recommended_action": f.recommended_action,
            }
            for f in findings
        ],
        "narrative": None,
        "prepared_by": "AI Assistant",
    }
